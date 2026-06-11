"""
Zero-Token Word (.docx) Assembler.

Reads a .docx template and injects user content into paragraphs identified
by para_index (from the docx parser's output).
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

from docx import Document

from app.schemas.mapping import TemplateMapping


def assemble_docx(
    template_path: str | Path,
    mapping: TemplateMapping,
    user_sections: dict[str, str],
) -> Document:
    """Inject user content into a Word template.

    Args:
        template_path: Path to the .docx template file.
        mapping: Module structure from the docx parser.
        user_sections: dict of ``module_id → new_content``.

    Returns:
        A python-docx ``Document`` with content injected.
    """
    doc = Document(str(template_path))
    paragraphs = doc.paragraphs

    for module in mapping.modules:
        mod_id = module.id
        new_content = user_sections.get(mod_id, "")
        if not new_content:
            continue

        para_index = module.anchor_line
        if para_index is None or para_index >= len(paragraphs):
            continue

        para = paragraphs[para_index]

        # If the paragraph contains a placeholder pattern, replace it
        if module.anchor_type == "placeholder" and module.anchor_pattern:
            import re
            new_text = re.sub(module.anchor_pattern, new_content, para.text, flags=re.IGNORECASE)
        else:
            # For heading-style matches, replace the entire paragraph text
            new_text = new_content

        # Clear existing runs and set new text
        for run in para.runs:
            run.text = ""
        if para.runs:
            para.runs[0].text = new_text
        else:
            para.add_run(new_text)

    return doc


def save_assembled_docx(
    doc: Document, output_path: str | Path
) -> Path:
    """Persist the assembled document to disk."""
    out = Path(output_path)
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))
    return out
