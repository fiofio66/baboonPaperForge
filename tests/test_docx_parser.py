"""
Test Word (.docx) template parser.

Generates real .docx files via python-docx for integration testing.
"""

import tempfile
from pathlib import Path

import pytest
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from app.services.docx_parser import parse_docx_template


@pytest.fixture
def sample_docx_path(tmp_path: Path) -> Path:
    """Create a minimal .docx template with headings and placeholders."""
    doc = Document()

    # Title
    title = doc.add_paragraph("{{TITLE}}")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.runs[0].font.size = Pt(16)

    # Author
    author = doc.add_paragraph("{{AUTHOR}}")
    author.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Abstract heading (Heading 1 style)
    h = doc.add_heading("Abstract", level=1)
    doc.add_paragraph("{{ABSTRACT}}")

    # Keywords
    doc.add_paragraph("{{KEYWORDS}}")

    # Sections with Heading 1 style
    doc.add_heading("Introduction", level=1)
    doc.add_paragraph("[[INTRODUCTION]] — replace with your intro text.")

    doc.add_heading("Related Work", level=1)
    doc.add_paragraph("[[RELATED WORK]]")

    doc.add_heading("Methodology", level=1)
    doc.add_paragraph("[[METHODS]]")

    # Subsection with Heading 2
    doc.add_heading("Dataset", level=2)
    doc.add_paragraph("Describe your dataset here.")

    doc.add_heading("Model Architecture", level=2)
    doc.add_paragraph("Describe the model.")

    doc.add_heading("Experiments", level=1)
    doc.add_paragraph("[[EXPERIMENTS]]")

    doc.add_heading("Conclusion", level=1)
    doc.add_paragraph("[[CONCLUSION]]")

    doc.add_heading("References", level=1)
    doc.add_paragraph("[[REFERENCES]]")

    p = tmp_path / "template.docx"
    doc.save(str(p))
    return p


def test_docx_parser_finds_placeholders(sample_docx_path: Path):
    modules = parse_docx_template(sample_docx_path)

    types = {m["type"] for m in modules}
    assert "title" in types
    assert "author" in types
    assert "abstract" in types
    assert "keywords" in types


def test_docx_parser_finds_headings(sample_docx_path: Path):
    modules = parse_docx_template(sample_docx_path)

    sections = [m for m in modules if m["type"] == "section"]
    labels = {m["label"] for m in sections}
    assert "Introduction" in labels
    assert "Methodology" in labels
    assert "Conclusion" in labels


def test_docx_parser_heading_levels(sample_docx_path: Path):
    modules = parse_docx_template(sample_docx_path)

    heading1 = [m for m in modules if m.get("level") == 1]
    heading2 = [m for m in modules if m.get("level") == 2]

    assert len(heading1) >= 5  # Abstract, Intro, Related Work, Methodology, Experiments, Conclusion, References
    assert len(heading2) >= 2  # Dataset, Model Architecture


def test_docx_parser_finds_bibliography(sample_docx_path: Path):
    modules = parse_docx_template(sample_docx_path)

    bibs = [m for m in modules if m["type"] == "bibliography"]
    assert len(bibs) >= 1


def test_docx_parser_empty_file(tmp_path: Path):
    """Empty .docx should return at least a minimal fallback module."""
    doc = Document()
    p = tmp_path / "empty.docx"
    doc.save(str(p))

    modules = parse_docx_template(p)
    # Fallback: at least one module (title placeholder)
    assert len(modules) >= 1
    assert modules[0]["type"] == "title"


def test_docx_parser_all_have_anchor_line(sample_docx_path: Path):
    modules = parse_docx_template(sample_docx_path)
    for m in modules:
        assert "anchor_line" in m, f"Missing anchor_line in {m['id']}"
