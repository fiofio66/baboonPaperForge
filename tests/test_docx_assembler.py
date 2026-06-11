"""
Test Zero-Token Word (.docx) Assembler.
"""

from pathlib import Path

import pytest
from docx import Document

from app.schemas.mapping import ModuleMapping, TemplateMapping
from app.services.docx_assembler import assemble_docx


@pytest.fixture
def sample_docx_path(tmp_path: Path) -> Path:
    doc = Document()
    doc.add_paragraph("{{TITLE}}")
    doc.add_paragraph("{{AUTHOR}}")
    doc.add_heading("Abstract", level=1)
    doc.add_paragraph("{{ABSTRACT}}")
    doc.add_heading("Introduction", level=1)
    doc.add_paragraph("[[INTRODUCTION]] — placeholder text.")
    doc.add_heading("Conclusion", level=1)
    doc.add_paragraph("[[CONCLUSION]]")

    p = tmp_path / "template.docx"
    doc.save(str(p))
    return p


@pytest.fixture
def sample_mapping() -> TemplateMapping:
    return TemplateMapping(
        template_path="/fake/template.docx",
        format="docx",
        modules=[
            ModuleMapping(id="mod-001", type="title", label="Title",
                          anchor_line=0, anchor_type="placeholder",
                          anchor_pattern=r"\{\{\s*TITLE\s*\}\}", content="{{TITLE}}"),
            ModuleMapping(id="mod-002", type="author", label="Author",
                          anchor_line=1, anchor_type="placeholder",
                          anchor_pattern=r"\{\{\s*AUTHOR\s*\}\}", content="{{AUTHOR}}"),
            ModuleMapping(id="mod-003", type="abstract", label="Abstract",
                          anchor_line=3, anchor_type="placeholder",
                          anchor_pattern=r"\{\{\s*ABSTRACT\s*\}\}", content="{{ABSTRACT}}"),
            ModuleMapping(id="mod-004", type="section", label="Introduction",
                          anchor_line=5, anchor_type="placeholder",
                          anchor_pattern=r"\[\[\s*INTRODUCTION\s*\]\]",
                          content="[[INTRODUCTION]] — placeholder text."),
            ModuleMapping(id="mod-005", type="section", label="Conclusion",
                          anchor_line=7, anchor_type="placeholder",
                          anchor_pattern=r"\[\[\s*CONCLUSION\s*\]\]",
                          content="[[CONCLUSION]]"),
        ],
    )


def test_assemble_docx_title_injection(sample_docx_path: Path, sample_mapping: TemplateMapping):
    user = {"mod-001": "My Custom Paper Title"}
    doc = assemble_docx(sample_docx_path, sample_mapping, user)
    assert "My Custom Paper Title" in doc.paragraphs[0].text


def test_assemble_docx_abstract_injection(sample_docx_path: Path, sample_mapping: TemplateMapping):
    user = {"mod-003": "Real abstract text here."}
    doc = assemble_docx(sample_docx_path, sample_mapping, user)
    assert "Real abstract text here." in doc.paragraphs[3].text


def test_assemble_docx_section_injection(sample_docx_path: Path, sample_mapping: TemplateMapping):
    user = {"mod-004": "This is the actual introduction."}
    doc = assemble_docx(sample_docx_path, sample_mapping, user)
    para_text = doc.paragraphs[5].text
    assert "This is the actual introduction." in para_text
    assert "[[INTRODUCTION]]" not in para_text  # placeholder tag replaced


def test_assemble_docx_no_llm_usage():
    import inspect
    src = inspect.getsource(assemble_docx)
    forbidden = ["openai", "anthropic", "langchain", "llm", "chat", "completion"]
    src_lower = src.lower()
    for word in forbidden:
        assert word not in src_lower, f"Assembler must not call {word}"
